# -*- coding: utf-8 -*-
"""
Baan Sao Yong Hin - quantity take-off from the BIM model -> BOQ (4 categories) -> CPM plan -> S-curve -> payments.

Reads the element list built by make_model.py (same geometry as the SketchUp script and the preview) and writes:
  BaanSaoYongHin_BOQ_Plan.xlsx   สรุป (ปร.5) / BOQ (ปร.4, สูตร) / ถอดปริมาณ / แผนงาน (Gantt) / S-Curve / งวดงาน / ข้อสมมติ
  tracker/plan.js                plan + BOQ + 4D links for the construction tracker web app

Quantities come from the model (volumes, areas, lengths, counts per building). Unit prices are 2026 Thai
market ESTIMATES (material + labour, no Factor F) in the same basis as House BIM Studio TYPE03; change them
in the BOQ sheet. Nothing here replaces a designed and priced tender BOQ.
Run:  python3 boq.py      (needs openpyxl; LibreOffice recalculation is optional)
"""
import datetime as dt
import importlib.util
import json
import math
import os
import re
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location('mm', os.path.join(HERE, 'make_model.py'))
mm = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(mm)
E = mm.E

PROJECT = 'บ้านเสายงหิน (Baan Sao Yong Hin) - กลุ่มอาคารไม้ชั้นเดียว 3 หลัง + งานภายนอก'
START = dt.date(2026, 11, 2)                       # ASSUMED start (Monday, dry season)
FACTOR_F = 1.2846                                   # same as TYPE03 - replace with the Comptroller-General table value
FT3 = 35.3147                                       # m3 -> cu.ft (ลบ.ฟ.)
HOLIDAYS = {dt.date(2026, 12, 5), dt.date(2026, 12, 7), dt.date(2026, 12, 10), dt.date(2026, 12, 31), dt.date(2027, 1, 1),
            dt.date(2027, 4, 6), dt.date(2027, 4, 13), dt.date(2027, 4, 14), dt.date(2027, 4, 15)}

# ------------------------------------------------------------------ geometry helpers
def bar_len(p):
    return math.dist(p[1], p[2])


def poly_area3(pts):
    s = [0.0, 0.0, 0.0]
    for i, a in enumerate(pts):
        b = pts[(i + 1) % len(pts)]
        s[0] += a[1] * b[2] - a[2] * b[1]; s[1] += a[2] * b[0] - a[0] * b[2]; s[2] += a[0] * b[1] - a[1] * b[0]
    return math.sqrt(sum(v * v for v in s)) / 2


def vol(p):
    if p[0] == 'box':
        return (p[2] - p[1]) * (p[4] - p[3]) * (p[6] - p[5])
    if p[0] == 'bar':
        return bar_len(p) * p[3] * p[4]
    return poly_area3(p[1]) * p[3]


def centroid(p):
    if p[0] == 'box':
        return ((p[1] + p[2]) / 2, (p[3] + p[4]) / 2)
    if p[0] == 'bar':
        return ((p[1][0] + p[2][0]) / 2, (p[1][1] + p[2][1]) / 2)
    return (sum(q[0] for q in p[1]) / len(p[1]), sum(q[1] for q in p[1]) / len(p[1]))


def locate(xy):
    x, y = xy
    if -1.3 <= x <= 13.2 and 3.0 <= y <= 12.6:
        return 'A'
    if 20.0 <= x <= 32.5 and -2.6 <= y <= 12.5:
        return 'C'
    dx, dy = x - mm.D_F.o[0], y - mm.D_F.o[1]
    a = dx * mm.D_F.u[0] + dy * mm.D_F.u[1]
    b = dx * mm.D_F.v[0] + dy * mm.D_F.v[1]
    if -1.3 <= a <= 12.5 and -1.5 <= b <= 11.5:
        return 'D'
    return 'S'


def elem_bld(e):
    cs = [centroid(p) for p in e[5]]
    return locate((sum(c[0] for c in cs) / len(cs), sum(c[1] for c in cs) / len(cs)))


BLD = {'A': 'อาคาร A', 'C': 'อาคาร C', 'D': 'อาคาร D', 'S': 'ภายนอก'}

# ------------------------------------------------------------------ activities (CPM)  code, name, duration (workdays), preds [(code, type, lag)]
ACTS = [
    ('PRE', 'งานเตรียมการ วางผัง รั้ว/โรงเก็บวัสดุชั่วคราว', 6, []),
    ('F-A', 'ขุดดิน ฐานราก ตอม่อ คานคอดิน อาคาร A', 12, [('PRE', 'FS', 0)]),
    ('F-C', 'ขุดดิน ฐานราก ตอม่อ คานคอดิน อาคาร C', 12, [('F-A', 'SS', 6)]),
    ('F-D', 'ฐานราก ตอม่อ พื้นคอนกรีตโรงจอดรถ อาคาร D', 10, [('F-C', 'SS', 6)]),
    ('F-S', 'ฐานหินซุ้มเสา ตอม่อชานไม้', 5, [('F-D', 'FS', 0)]),
    ('UG-A', 'ท่อประปา/ระบาย/ท่อร้อยสายใต้พื้น อาคาร A', 3, [('F-A', 'FS', 0)]),
    ('UG-C', 'ท่อประปา/ระบาย/ท่อร้อยสายใต้พื้น อาคาร C', 3, [('F-C', 'FS', 0)]),
    ('SL-A', 'ถมทราย เทพื้นคอนกรีต อาคาร A', 4, [('UG-A', 'FS', 0)]),
    ('SL-C', 'ถมทราย เทพื้นคอนกรีต อาคาร C', 4, [('UG-C', 'FS', 0)]),
    ('TS-A', 'โครงสร้างไม้ เสา คาน อาคาร A', 10, [('SL-A', 'FS', 0)]),
    ('TS-C', 'โครงสร้างไม้ เสา คาน อาคาร C', 10, [('SL-C', 'FS', 0), ('TS-A', 'FS', 0)]),
    ('TS-D', 'โครงสร้างไม้ เสา คาน อาคาร D', 7, [('F-D', 'FS', 0), ('TS-C', 'FS', 0)]),
    ('RF-A', 'โครงหลังคา มุงหลังคาลอน อาคาร A', 9, [('TS-A', 'FS', 0)]),
    ('RF-C', 'โครงหลังคา มุงหลังคาลอน อาคาร C', 9, [('TS-C', 'FS', 0), ('RF-A', 'FS', 0)]),
    ('RF-D', 'โครงหลังคา หลังคายก มุงหลังคา อาคาร D', 8, [('TS-D', 'FS', 0), ('RF-C', 'FS', 0)]),
    ('PG', 'ซุ้มเสาไม้ (pergola) บนฐานหิน', 6, [('F-S', 'FS', 0), ('TS-D', 'FS', 0)]),
    ('WL-A', 'ผนังโครงไม้ กรุไม้ ผนังก่อห้องน้ำ อาคาร A', 14, [('RF-A', 'FS', 0)]),
    ('WL-C', 'ผนังโครงไม้ กรุไม้ อาคาร C', 14, [('RF-C', 'FS', 0), ('WL-A', 'SS', 7)]),
    ('WL-D', 'ห้องเก็บของ ระแนงไม้ กำแพงกาเบียน งานระบบ อาคาร D', 10, [('RF-D', 'FS', 0)]),
    ('MR-A', 'เดินสายไฟ-ท่อน้ำในผนัง (rough-in) อาคาร A', 5, [('WL-A', 'SS', 5)]),
    ('MR-C', 'เดินสายไฟ-ท่อน้ำในผนัง (rough-in) อาคาร C', 5, [('WL-C', 'SS', 5)]),
    ('CL-A', 'ฝ้าไม้ไผ่สาน/ฝ้ากันชื้น อาคาร A', 6, [('WL-A', 'FS', 0), ('MR-A', 'FS', 0)]),
    ('CL-C', 'ฝ้าไม้ อาคาร C', 6, [('WL-C', 'FS', 0), ('MR-C', 'FS', 0)]),
    ('FL-A', 'พื้นขัดมัน กันซึม ปูกระเบื้อง อาคาร A', 6, [('CL-A', 'FS', 0)]),
    ('FL-C', 'พื้นขัดมัน กันซึม ปูกระเบื้อง อาคาร C', 6, [('CL-C', 'FS', 0)]),
    ('DW-A', 'ประตู-หน้าต่างไม้เก่า บานเฟี้ยมอลูมิเนียม อาคาร A', 7, [('WL-A', 'FS', 0)]),
    ('DW-C', 'ประตู-หน้าต่างไม้เก่า บานเลื่อน อาคาร C', 7, [('WL-C', 'FS', 0)]),
    ('PT-A', 'ทาสี เคลือบไม้ อาคาร A', 5, [('FL-A', 'FS', 0), ('DW-A', 'FS', 0)]),
    ('PT-C', 'ทาสี เคลือบไม้ อาคาร C', 5, [('FL-C', 'FS', 0), ('DW-C', 'FS', 0)]),
    ('FX-A', 'ติดตั้งสุขภัณฑ์ ดวงโคม แอร์ เครื่องทำน้ำอุ่น อาคาร A', 4, [('PT-A', 'FS', 0)]),
    ('FX-C', 'ติดตั้งสุขภัณฑ์ ดวงโคม แอร์ เครื่องทำน้ำอุ่น อาคาร C', 4, [('PT-C', 'FS', 0)]),
    ('DK', 'ชานไม้ ระเบียง ทางเดินไม้ ชานอ่างอาบน้ำ', 14, [('TS-A', 'FS', 0), ('PG', 'FS', 0)]),
    ('UT', 'งานระบบภายนอก: ถังบำบัด บ่อซึม ถังน้ำ ปั๊ม สายเมนใต้ดิน มิเตอร์', 10, [('F-D', 'FS', 0)]),
    ('EX', 'งานภายนอก: ลานคอนกรีต หินกรวด รางหินรับน้ำฝน ไฟสนาม', 6, [('DK', 'FS', 0), ('UT', 'FS', 0)]),
    ('TC', 'ทดสอบระบบไฟฟ้า-ประปา ขอใช้ไฟ/น้ำ', 5, [('FX-A', 'FS', 0), ('FX-C', 'FS', 0), ('WL-D', 'FS', 0), ('UT', 'FS', 0)]),
    ('HO', 'ทำความสะอาด ตรวจรับ ส่งมอบ แบบ As-built', 4, [('TC', 'FS', 0), ('EX', 'FS', 0)]),
]
ACT_IDS = [a[0] for a in ACTS]
FALLBACK = {'F-S': 'F-S', 'SL-D': 'F-D', 'SL-S': 'EX', 'TS-S': 'PG', 'RF-S': 'PG', 'WL-S': 'EX', 'CL-D': 'WL-D', 'FL-D': 'F-D', 'FL-S': 'EX',
            'DW-D': 'WL-D', 'PT-D': 'WL-D', 'PT-S': 'DK', 'FX-D': 'WL-D', 'FX-S': 'EX', 'MR-D': 'WL-D', 'MR-S': 'EX', 'UG-D': 'UT',
            'UG-S': 'UT', 'F-D': 'F-D'}


def act_of(tmpl, b):
    code = tmpl.format(b=b)
    if code in ACT_IDS:
        return code
    if code in FALLBACK:
        return FALLBACK[code]
    raise KeyError(code)


# ------------------------------------------------------------------ preliminary design (design_calc.py), if present
_DJ = os.path.join(HERE, 'design', 'design_results.json')
DESIGN = json.load(open(_DJ, encoding='utf-8')) if os.path.exists(_DJ) else {}


def _feeder(which):
    el = DESIGN.get('electrical')
    if not el:
        return None
    if which == 'main':
        m = el['main']
        return dict(cable=m['cable'], conduit='', mm2=m['mm2'], cores=4)
    f = {x['to']: x for x in el['feeders']}['CU-' + which]
    return dict(cable=f['cable'], conduit=f['conduit'], mm2=f['mm2'], cores=4 if f['three_phase'] else 2, ground=f['ground'])


def _cable(which, default):
    f = _feeder(which)
    if not f:
        return default
    txt = f['cable'].replace('mm2', 'ตร.มม.').replace(' in ', ' ใน ').replace(' mm', ' มม.')
    return txt + (' ใน ' + f['conduit'].replace(' mm', ' มม.') if f['conduit'] else '')


def _cable_rate(which, default):
    """NYY cable + HDPE conduit, material per metre (ESTIMATE: ~4.6 baht per core-mm2 per m + 85 baht conduit/fittings)"""
    f = _feeder(which)
    if not f:
        return default
    return round(4.6 * (f['cores'] * f['mm2'] + f.get('ground', 0)) + 85, -1)


# ------------------------------------------------------------------ BOQ items: key -> (cat, group, description, unit, mat, lab, waste%, activity template)
S, A, P, EL = 'S', 'A', 'P', 'E'
CATS = [(S, 'หมวดที่ 1 งานโครงสร้าง'), (A, 'หมวดที่ 2 งานสถาปัตยกรรม'), (P, 'หมวดที่ 3 งานระบบสุขาภิบาลและประปา'), (EL, 'หมวดที่ 4 งานระบบไฟฟ้าและแสงสว่าง')]
ITEMS = {
    # ---- structure
    'setout': (S, '1.1 งานเตรียมการ', 'สำรวจ วางผัง ตั้งหมุดแนวเสา ระดับ', 'งาน', 0, 12000, 0, 'PRE'),
    'temp': (S, '1.1 งานเตรียมการ', 'งานชั่วคราว: รั้ว ป้าย โรงเก็บวัสดุ ไฟ-น้ำชั่วคราว', 'งาน', 25000, 15000, 0, 'PRE'),
    'clear': (S, '1.1 งานเตรียมการ', 'ถางปรับพื้นที่ก่อสร้าง (เก็บต้นไม้เดิม 2 ต้น)', 'ตร.ม.', 0, 25, 0, 'PRE'),
    'exc': (S, '1.2 งานดินและฐานราก', 'ขุดดินฐานราก คานคอดิน ตอม่อชาน', 'ลบ.ม.', 0, 150, 0, 'F-{b}'),
    'lean': (S, '1.2 งานดินและฐานราก', 'ทรายหยาบรองพื้น + คอนกรีตหยาบ 1:3:5 หนา 0.05', 'ลบ.ม.', 2150, 300, 5, 'F-{b}'),
    'c_ft': (S, '1.2 งานดินและฐานราก', "คอนกรีตฐานราก fc' 240 ksc", 'ลบ.ม.', 2350, 300, 3, 'F-{b}'),
    'c_st': (S, '1.2 งานดินและฐานราก', "คอนกรีตตอม่อ fc' 240 ksc", 'ลบ.ม.', 2350, 350, 3, 'F-{b}'),
    'c_gb': (S, '1.2 งานดินและฐานราก', "คอนกรีตคานคอดิน fc' 240 ksc", 'ลบ.ม.', 2350, 300, 3, 'F-{b}'),
    'c_pier': (S, '1.2 งานดินและฐานราก', 'คอนกรีตตอม่อชานไม้ 0.20x0.20 + ฐานเล็ก', 'ลบ.ม.', 2350, 400, 3, 'F-{b}'),
    'fw': (S, '1.2 งานดินและฐานราก', 'ไม้แบบ ฐานราก ตอม่อ คานคอดิน', 'ตร.ม.', 160, 200, 0, 'F-{b}'),
    'rebar': (S, '1.2 งานดินและฐานราก', 'เหล็กเสริม SD40 DB12 / SR24 RB6, RB9 (ฐานราก ตอม่อ คานคอดิน ตามรายการคำนวณ design_calc.py)', 'กก.', 25, 5, 5, 'F-{b}'),
    'boulder': (S, '1.2 งานดินและฐานราก', 'ฐานหินธรรมชาติใต้เสาซุ้ม + เดือยเหล็ก/แผ่นยึด', 'ก้อน', 1500, 500, 0, 'F-{b}'),
    'backfill': (S, '1.2 งานดินและฐานราก', 'ถมดินคืนบดอัด', 'ลบ.ม.', 0, 90, 0, 'F-{b}'),
    'sand_fill': (S, '1.3 งานพื้น', 'ทรายถมบดอัดใต้พื้น หนา 0.10', 'ลบ.ม.', 500, 100, 5, 'SL-{b}'),
    'pe': (S, '1.3 งานพื้น', 'แผ่นพลาสติก PE กันความชื้นใต้พื้น', 'ตร.ม.', 12, 3, 10, 'SL-{b}'),
    'mesh': (S, '1.3 งานพื้น', 'ตะแกรงเหล็ก wire mesh 4 มม. @0.20', 'ตร.ม.', 65, 15, 10, 'SL-{b}'),
    'c_sl': (S, '1.3 งานพื้น', "คอนกรีตพื้นวางบนดิน fc' 240 ksc", 'ลบ.ม.', 2350, 300, 3, 'SL-{b}'),
    'fw_sl': (S, '1.3 งานพื้น', 'ไม้แบบขอบพื้น', 'ตร.ม.', 160, 200, 0, 'SL-{b}'),
    'tm_post': (S, '1.4 งานโครงสร้างไม้', 'เสาไม้เนื้อแข็งเก่า (เต็ง/รัง) ไสตกแต่ง', 'ลบ.ฟ.', 750, 250, 10, 'TS-{b}'),
    'tm_beam': (S, '1.4 งานโครงสร้างไม้', 'คาน/อะเส/ตงไม้เนื้อแข็งเก่า', 'ลบ.ฟ.', 700, 250, 10, 'TS-{b}'),
    'tm_roof': (S, '1.4 งานโครงสร้างไม้', 'จันทัน แป อกไก่ ไม้เนื้อแข็งเก่า', 'ลบ.ฟ.', 650, 220, 10, 'RF-{b}'),
    'tm_steel': (S, '1.4 งานโครงสร้างไม้', 'เหล็กประกับ สลักเกลียว พุก (25 กก./ลบ.ม.ไม้)', 'กก.', 60, 20, 5, 'TS-{b}'),
    'tm_treat': (S, '1.4 งานโครงสร้างไม้', 'อาบน้ำยากันปลวก-กันเชื้อรา ไม้โครงสร้าง', 'ลบ.ฟ.', 30, 10, 0, 'TS-{b}'),
    'wl_brace': (S, '1.4 งานโครงสร้างไม้', 'ค้ำยันทแยงในผนัง 50x100 (ต้านแรงลมด้านข้าง) + สลัก M12 2 ตัว/ปลาย', 'ชุด', 450, 250, 0, 'WL-{b}'),
    'gabion': (S, '1.5 งานกำแพง', 'กำแพงกาเบียน ตะแกรงเหล็กชุบสังกะสี + หินแม่น้ำ', 'ลบ.ม.', 1400, 500, 5, 'WL-{b}'),
    # ---- architecture
    'wl_frame': (A, '2.1 งานผนัง', 'โครงผนังไม้ (เคร่า) 2"x3" @0.40', 'ตร.ม.', 220, 120, 5, 'WL-{b}'),
    'cl_board': (A, '2.1 งานผนัง', 'ไม้เก่ากรุผนังแนวนอน (ต่อด้าน)', 'ตร.ม.', 450, 180, 10, 'WL-{b}'),
    'cl_vboard': (A, '2.1 งานผนัง', 'ไม้กรุแนวตั้งหน้าจั่ว/เหนือบานกระจก รวมโครง', 'ตร.ม.', 480, 200, 10, 'WL-{b}'),
    'wl_aac': (A, '2.1 งานผนัง', 'ผนังอิฐมวลเบา หนา 0.10-0.15 (ห้องน้ำ ผนังสีขาว)', 'ตร.ม.', 230, 110, 3, 'WL-{b}'),
    'plaster': (A, '2.1 งานผนัง', 'ฉาบปูนผนังอิฐมวลเบา (2 ด้าน)', 'ตร.ม.', 60, 90, 3, 'WL-{b}'),
    'screen': (A, '2.1 งานผนัง', 'ระแนงไม้เก่า (แนวตั้ง/แนวนอน) รวมโครง', 'ตร.ม.', 550, 250, 10, 'WL-{b}'),
    'fl_polish': (A, '2.2 งานพื้นผิว', 'พื้นปูนขัดมันผสมน้ำยากันซึม + เคลือบ sealer', 'ตร.ม.', 120, 150, 5, 'FL-{b}'),
    'fl_trowel': (A, '2.2 งานพื้นผิว', 'พื้นคอนกรีตขัดหยาบ (โรงจอดรถ)', 'ตร.ม.', 25, 60, 0, 'FL-{b}'),
    'wp': (A, '2.2 งานพื้นผิว', 'กันซึมห้องน้ำ (พื้น + ผนังสูง 0.30)', 'ตร.ม.', 180, 60, 5, 'FL-{b}'),
    'fl_tile': (A, '2.2 งานพื้นผิว', 'กระเบื้องพื้นห้องน้ำ 30x30 กันลื่น', 'ตร.ม.', 450, 180, 5, 'FL-{b}'),
    'wt_tile': (A, '2.2 งานพื้นผิว', 'กระเบื้องผนังส่วนอาบน้ำ สูง 1.80', 'ตร.ม.', 450, 200, 5, 'FL-{b}'),
    'deck': (A, '2.2 งานพื้นผิว', 'พื้นชานไม้เก่า หนา 1" บนตงไม้ (ชาน ระเบียง ทางเดิน)', 'ตร.ม.', 1800, 450, 10, 'DK'),
    'ceil_bamboo': (A, '2.3 งานฝ้าเพดาน', 'ฝ้าไม้ไผ่สาน ตามแนวลาดหลังคา บนโครงไม้', 'ตร.ม.', 380, 180, 5, 'CL-{b}'),
    'ceil_timber': (A, '2.3 งานฝ้าเพดาน', 'ฝ้าไม้ตีเกล็ด/ลิ้น บนโครงไม้', 'ตร.ม.', 420, 180, 5, 'CL-{b}'),
    'ceil_wr': (A, '2.3 งานฝ้าเพดาน', 'ฝ้าแผ่นกันชื้น ห้องน้ำ', 'ตร.ม.', 280, 130, 5, 'CL-{b}'),
    'rf_sheet': (A, '2.4 งานหลังคา', 'หลังคากระเบื้องลอนไฟเบอร์ซีเมนต์ สีเทา + ขอยึด', 'ตร.ม.', 190, 70, 10, 'RF-{b}'),
    'rf_ridge': (A, '2.4 งานหลังคา', 'ครอบสันหลังคา', 'ม.', 120, 40, 5, 'RF-{b}'),
    'rf_flash': (A, '2.4 งานหลังคา', 'แผ่นปิดรอยต่อ (flashing) ช่องแสงหลังคายก/หลังคาซ้อน', 'ม.', 180, 60, 5, 'RF-{b}'),
    'rf_fascia': (A, '2.4 งานหลังคา', 'ไม้ปั้นลม/เชิงชาย', 'ม.', 220, 80, 5, 'RF-{b}'),
    'dr_reclaim': (A, '2.5 งานประตู-หน้าต่าง', 'ประตูไม้เก่า (ตามตาราง D) ซ่อม-ทำวงกบ-อุปกรณ์-ติดตั้ง', 'ชุด', 2500, 1500, 0, 'DW-{b}'),
    'wn_reclaim_l': (A, '2.5 งานประตู-หน้าต่าง', 'หน้าต่างไม้เก่าชุดใหญ่ (>1.5 ตร.ม.) ซ่อม-วงกบ-กระจกช่องแสง-ติดตั้ง', 'ชุด', 2800, 1500, 0, 'DW-{b}'),
    'wn_reclaim_s': (A, '2.5 งานประตู-หน้าต่าง', 'หน้าต่างไม้เก่า/ช่องระบาย ชุดเล็ก ซ่อม-วงกบ-ติดตั้ง', 'ชุด', 1600, 900, 0, 'DW-{b}'),
    'al_bifold': (A, '2.5 งานประตู-หน้าต่าง', 'ประตูบานเฟี้ยมอลูมิเนียมสีเทา กระจกเทมเปอร์ 6 มม.', 'ตร.ม.', 4200, 600, 0, 'DW-{b}'),
    'al_slide': (A, '2.5 งานประตู-หน้าต่าง', 'ประตูบานเลื่อนอลูมิเนียม กระจกเทมเปอร์ 6 มม.', 'ตร.ม.', 3200, 500, 0, 'DW-{b}'),
    'stain_ext': (A, '2.6 งานสีและเคลือบผิว', 'น้ำมันรักษาเนื้อไม้/สีย้อมไม้ ผนัง ระแนง โครงสร้างไม้', 'ตร.ม.', 40, 40, 5, 'PT-{b}'),
    'paint_wall': (A, '2.6 งานสีและเคลือบผิว', 'สีอะคริลิกผนังปูน รองพื้น + ทับหน้า 2 เที่ยว', 'ตร.ม.', 45, 42, 5, 'PT-{b}'),
    'deck_oil': (A, '2.6 งานสีและเคลือบผิว', 'น้ำมันเคลือบพื้นชานไม้', 'ตร.ม.', 45, 35, 5, 'DK'),
    'terrace': (A, '2.7 งานภายนอก', 'ลานคอนกรีต หนา 0.10 เสริมตะแกรง ขัดหยาบ', 'ตร.ม.', 450, 200, 3, 'EX'),
    'gravel': (A, '2.7 งานภายนอก', 'ลานหินกรวด หนา 0.10 บนแผ่นกันวัชพืช (ทางรถ)', 'ตร.ม.', 150, 40, 5, 'EX'),
    'drip': (A, '2.7 งานภายนอก', 'รางหินกรวดรับน้ำฝนใต้ชายคา กว้าง 0.40', 'ม.', 120, 60, 5, 'EX'),
    'steps': (A, '2.7 งานภายนอก', 'ขั้นบันไดไม้หมอน + หินรองทางขึ้น', 'ชุด', 3500, 1500, 0, 'EX'),
    'bench': (A, '2.7 งานภายนอก', 'ม้านั่ง/กล่องที่นั่งไม้ บิลท์อิน', 'ชุด', 6000, 2500, 0, 'EX'),
    'clean': (A, '2.8 งานส่งมอบ', 'ทำความสะอาด ขนย้ายเศษวัสดุ', 'งาน', 3000, 12000, 0, 'HO'),
    'asbuilt': (A, '2.8 งานส่งมอบ', 'แบบ As-built + คู่มือบำรุงรักษา', 'งาน', 0, 15000, 0, 'HO'),
    # ---- sanitary / plumbing
    'san_wc': (P, '3.1 สุขภัณฑ์และอุปกรณ์', 'ชักโครก + สายฉีดชำระ', 'ชุด', 6000, 800, 0, 'FX-{b}'),
    'san_lav': (P, '3.1 สุขภัณฑ์และอุปกรณ์', 'อ่างล้างหน้า + ก๊อก + ท่อน้ำทิ้ง', 'ชุด', 3500, 600, 0, 'FX-{b}'),
    'san_shower': (P, '3.1 สุขภัณฑ์และอุปกรณ์', 'ชุดฝักบัว rain shower + วาล์วผสม', 'ชุด', 6500, 900, 0, 'FX-{b}'),
    'san_sink': (P, '3.1 สุขภัณฑ์และอุปกรณ์', 'เคาน์เตอร์ pantry 2.0 ม. + ซิงค์สแตนเลส + ก๊อก', 'ชุด', 9000, 2000, 0, 'FX-{b}'),
    'san_tub': (P, '3.1 สุขภัณฑ์และอุปกรณ์', 'อ่างแช่กลางแจ้งทรงกลม dia 1.10 + ก๊อก + ท่อทิ้ง', 'ชุด', 25000, 3000, 0, 'FX-{b}'),
    'p_fd': (P, '3.1 สุขภัณฑ์และอุปกรณ์', 'ตะแกรงระบายน้ำพื้น 4" สแตนเลส กันกลิ่น', 'ชุด', 400, 200, 0, 'FL-{b}'),
    'p_yard': (P, '3.1 สุขภัณฑ์และอุปกรณ์', 'ก๊อกสนามบนเสา', 'จุด', 450, 300, 0, 'UT'),
    'p_meter': (P, '3.2 ระบบน้ำประปา', 'ขอติดตั้งมาตรวัดน้ำ 1/2" + ประตูน้ำ + เช็ควาล์ว (ค่าธรรมเนียมประมาณ)', 'งาน', 6500, 1500, 0, 'UT'),
    'p_tank': (P, '3.2 ระบบน้ำประปา', 'ถังเก็บน้ำ PE 2,000 ลิตร + แท่นคอนกรีต + ลูกลอย', 'ชุด', 7500, 1500, 0, 'UT'),
    'p_pump': (P, '3.2 ระบบน้ำประปา', 'ปั๊มน้ำอัตโนมัติแรงดันคงที่ 300 W + หลังคาคลุม', 'ชุด', 12000, 2000, 0, 'UT'),
    'pe25': (P, '3.2 ระบบน้ำประปา', 'ท่อ HDPE PN10 dia 25 มม. ฝังดิน รวมข้อต่อ', 'ม.', 35, 30, 10, 'UG-{b}'),
    'ppr25': (P, '3.2 ระบบน้ำประปา', 'ท่อ PPR PN20 dia 25 มม. น้ำเย็น รวมข้อต่อ', 'ม.', 55, 45, 10, 'UG-{b}'),
    'ppr20': (P, '3.2 ระบบน้ำประปา', 'ท่อ PPR PN20 dia 20 มม. น้ำเย็น รวมข้อต่อ', 'ม.', 45, 40, 10, 'MR-{b}'),
    'ppr20h': (P, '3.2 ระบบน้ำประปา', 'ท่อ PPR PN20 dia 20 มม. น้ำร้อน รวมข้อต่อ', 'ม.', 55, 45, 10, 'MR-{b}'),
    'pvc100': (P, '3.3 ระบบระบายน้ำ', 'ท่อ PVC ชั้น 8.5 dia 100 มม. โสโครก รวมข้อต่อ', 'ม.', 120, 60, 10, 'UG-{b}'),
    'pvc80': (P, '3.3 ระบบระบายน้ำ', 'ท่อ PVC ชั้น 8.5 dia 80 มม. รวมข้อต่อ', 'ม.', 85, 50, 10, 'UG-{b}'),
    'pvc55': (P, '3.3 ระบบระบายน้ำ', 'ท่อ PVC ชั้น 8.5 dia 55 มม. น้ำทิ้ง รวมข้อต่อ', 'ม.', 60, 45, 10, 'UG-{b}'),
    'p_ic': (P, '3.3 ระบบระบายน้ำ', 'บ่อพัก 0.50x0.50 สำเร็จรูป + ฝา', 'บ่อ', 1200, 600, 0, 'UT'),
    'p_gt': (P, '3.3 ระบบระบายน้ำ', 'บ่อดักไขมัน 30 ลิตร ฝังดิน', 'ชุด', 2500, 600, 0, 'UT'),
    'p_septic': (P, '3.3 ระบบระบายน้ำ', 'ถังบำบัดน้ำเสียสำเร็จรูป 1,600 ลิตร', 'ชุด', 14000, 3500, 0, 'UT'),
    'p_soak': (P, '3.3 ระบบระบายน้ำ', 'บ่อซึม dia 1.20 ลึก 1.80 (ท่อซีเมนต์ + หินกรวด)', 'บ่อ', 4500, 2500, 0, 'UT'),
    'p_test': (P, '3.4 ทดสอบ', 'ทดสอบแรงดันท่อประปา + ทดสอบการไหลท่อระบาย', 'งาน', 0, 3000, 0, 'TC'),
    # ---- electrical
    'e_meter': (EL, '4.1 ระบบไฟฟ้ากำลัง', 'ขอใช้ไฟ PEA มิเตอร์ %s + เสาคอนกรีต 6 ม. (ค่าธรรมเนียมประมาณ)' % (
        DESIGN['electrical']['main']['meter'].replace('3-phase 4-wire', '3 เฟส 4 สาย') if DESIGN.get('electrical') else '15(45)A 3 เฟส'), 'งาน', 22000, 6000, 0, 'UT'),
    'e_mdb': (EL, '4.1 ระบบไฟฟ้ากำลัง', 'ตู้ MDB load center 3 เฟส เมน %s 12 ช่อง' % (DESIGN['electrical']['main']['main_cb'] if DESIGN.get('electrical') else '3P 63A + RCBO'),
              'ชุด', 18000, 3000, 0, 'WL-{b}'),
    'e_cu': (EL, '4.1 ระบบไฟฟ้ากำลัง', 'ตู้ consumer unit 1 เฟส เมน + RCBO/RCD 30 mA 10 ช่อง', 'ชุด', 6500, 1500, 0, 'MR-{b}'),
    'e_cu3': (EL, '4.1 ระบบไฟฟ้ากำลัง', 'ตู้ consumer unit 3 เฟส 4 สาย เมน 3P + RCBO/RCD 30 mA 12 ช่อง', 'ชุด', 11000, 2000, 0, 'MR-{b}'),
    'e_fd_main': (EL, '4.1 ระบบไฟฟ้ากำลัง', 'สายเมน %s ฝังดิน 0.65 ม. รวมขุด-กลบ' % _cable('main', 'NYY 4x25 ตร.มม. ใน HDPE 50 มม.'), 'ม.',
                  _cable_rate('main', 520), 120, 5, 'UG-{b}'),
    'e_fd_a': (EL, '4.1 ระบบไฟฟ้ากำลัง', 'สายป้อน %s ฝังดิน MDB -> CU-A' % _cable('A', 'NYY 2x16 + G 10 ตร.มม. ใน HDPE 40 มม.'), 'ม.',
               _cable_rate('A', 250), 90, 5, 'UG-{b}'),
    'e_fd_c': (EL, '4.1 ระบบไฟฟ้ากำลัง', 'สายป้อน %s ฝังดิน MDB -> CU-C' % _cable('C', 'NYY 2x25 + G 16 ตร.มม. ใน HDPE 40 มม.'), 'ม.',
               _cable_rate('C', 330), 100, 5, 'UG-{b}'),
    'e_fd_site': (EL, '4.1 ระบบไฟฟ้ากำลัง', 'สาย NYY 3x2.5 ตร.มม. ใน HDPE 25 มม. วงจรไฟสนาม', 'ม.', 95, 60, 5, 'UG-{b}'),
    'e_fd_pump': (EL, '4.1 ระบบไฟฟ้ากำลัง', 'สาย NYY 3x2.5 ตร.มม. ใน HDPE 25 มม. วงจรปั๊มน้ำ', 'ม.', 95, 60, 5, 'UG-{b}'),
    'e_sock': (EL, '4.2 เต้ารับและสวิตช์', 'จุดเต้ารับคู่มีกราวด์ 16A (สาย THW 2.5 ในท่อ)', 'จุด', 450, 300, 0, 'MR-{b}'),
    'e_sock_ip': (EL, '4.2 เต้ารับและสวิตช์', 'จุดเต้ารับกันน้ำ IP55 ภายนอก', 'จุด', 700, 350, 0, 'MR-{b}'),
    'e_switch': (EL, '4.2 เต้ารับและสวิตช์', 'จุดสวิตช์ไฟ 1-2 ทาง', 'จุด', 250, 200, 0, 'MR-{b}'),
    'e_lpt': (EL, '4.3 ระบบแสงสว่าง', 'จุดเดินสายดวงโคม (THW 2.5 + G ในท่อ เฉลี่ย 8 ม./จุด, ขนาดต่ำสุดตาม วสท.)', 'จุด', 380, 250, 0, 'MR-{b}'),
    'l_pend': (EL, '4.3 ระบบแสงสว่าง', 'โคมไฟห้อย โป๊ะกระดาษ LED 12 W', 'ชุด', 1800, 250, 0, 'FX-{b}'),
    'l_down': (EL, '4.3 ระบบแสงสว่าง', 'โคมดาวน์ไลท์ LED 9 W 3000K', 'ชุด', 450, 150, 0, 'FX-{b}'),
    'l_wall': (EL, '4.3 ระบบแสงสว่าง', 'โคมติดผนัง/เสา LED 7 W IP54', 'ชุด', 900, 200, 0, 'FX-{b}'),
    'l_batten': (EL, '4.3 ระบบแสงสว่าง', 'โคม LED batten 1.2 ม. 18 W IP65', 'ชุด', 650, 150, 0, 'FX-{b}'),
    'l_flood': (EL, '4.3 ระบบแสงสว่าง', 'โคมฟลัดไลท์ LED 30 W IP65 + photocell', 'ชุด', 1500, 300, 0, 'FX-{b}'),
    'l_bollard': (EL, '4.3 ระบบแสงสว่าง', 'โคมไฟสนามแบบเสาเตี้ย LED 5 W IP65', 'ชุด', 1600, 400, 0, 'FX-{b}'),
    'e_wh': (EL, '4.4 เครื่องใช้ไฟฟ้า', 'เครื่องทำน้ำอุ่น 4,500 W + วงจร RCBO 32A สาย THW 6', 'ชุด', 4500, 800, 0, 'FX-{b}'),
    'e_ac': (EL, '4.4 เครื่องใช้ไฟฟ้า', 'เครื่องปรับอากาศ 12,000 BTU inverter + ท่อ 4 ม. + วงจร 20A', 'ชุด', 16000, 3500, 0, 'FX-{b}'),
    'e_ac_cdu': (EL, '4.4 เครื่องใช้ไฟฟ้า', 'แท่นคอนกรีต/ขาแขวน คอยล์ร้อน', 'ชุด', 600, 300, 0, 'FX-{b}'),
    'e_earth': (EL, '4.5 ระบบสายดินและทดสอบ', 'หลักดินทองแดง 5/8" x 2.4 ม. + บ่อตรวจ', 'ชุด', 1200, 600, 0, 'UT'),
    'e_test': (EL, '4.5 ระบบสายดินและทดสอบ', 'ทดสอบฉนวน ค่าความต้านทานดิน RCD + ใบรับรอง', 'งาน', 0, 5000, 0, 'TC'),
}

# ------------------------------------------------------------------ take-off
Q = defaultdict(float)                    # (item, bld) -> qty
TRACE = []                                # element, ifc, tag, bld, item, qty, unit, basis
ELEM_ACT = {}                             # element name -> activity (4D link: first/main item)


def put(item, b, q, e=None, basis=''):
    if q <= 0:
        return
    Q[(item, b)] += q
    TRACE.append((e[1] if e else '-', e[0] if e else '-', e[4] if e else '-', b, item, round(q, 3), ITEMS[item][3], basis))
    if e is not None and e[1] not in ELEM_ACT:
        ELEM_ACT[e[1]] = act_of(ITEMS[item][7], b)


def mats(e):
    return {p[-1] for p in e[5]}


for k, q in (('setout', 1), ('temp', 1)):
    put(k, 'S', q, basis='lump sum')
foot = 0.0
for e in E:
    ifc, name, mark, lvl, tag, parts, at = e
    b = elem_bld(e)
    if name in ('A-SLAB', 'C-SLAB', 'D-SLAB'):
        foot += sum((p[2] - p[1]) * (p[4] - p[3]) if p[0] == 'box' else poly_area3(p[1]) for p in parts)
put('clear', 'S', round(foot * 1.6 + 400, 0), basis='building slabs x 1.6 + decks/pergola area ~400 m2')

for e in E:
    ifc, name, mark, lvl, tag, parts, at = e
    b = elem_bld(e)
    m = mats(e)
    v = sum(vol(p) for p in parts)
    # ---------------- structure
    if ifc == 'IfcFooting' and 'concrete' in m:
        p0 = parts[0]
        Bf = (p0[2] - p0[1]) if p0[0] == 'box' else math.dist(p0[1][0], p0[1][1])
        hf = (p0[6] - p0[5]) if p0[0] == 'box' else p0[3]
        pit = (Bf + 0.4) ** 2 * (0.95 + hf + 0.05)
        put('c_ft', b, v, e, 'model volume'); put('exc', b, pit, e, '(B+0.40)^2 x depth')
        put('lean', b, (Bf + 0.2) ** 2 * 0.10, e, '(B+0.20)^2 sand 0.05 + lean 0.05'); put('fw', b, 4 * Bf * hf, e, 'perimeter x h')
        put('rebar', b, at.get('Rebar_kg', v * 60), e, 'designed: ' + at['Rebar'] if 'Rebar' in at else '60 kg/m3')
        put('backfill', b, pit - v - 0.04 * 0.95 - (Bf + 0.2) ** 2 * 0.10, e, 'pit - concrete - stump - lean')
    elif ifc == 'IfcFooting':
        put('boulder', b, 1, e, 'count')
    elif ifc == 'IfcColumn' and mark == 'ST':
        h = parts[0][6] - parts[0][5] if parts[0][0] == 'box' else parts[0][3]
        put('c_st', b, v, e, 'model volume'); put('fw', b, 4 * 0.2 * h, e, '4 x 0.20 x h'); put('rebar', b, at.get('Rebar_kg', v * 140), e, 'designed: ' + at['Rebar'] if 'Rebar' in at else '140 kg/m3')
    elif ifc == 'IfcBeam' and mark == 'GB1':
        L = sum(max(p[2] - p[1], p[4] - p[3]) for p in parts)
        put('c_gb', b, v, e, 'model volume'); put('fw', b, 2 * 0.4 * L, e, '2 sides x 0.40 x L'); put('rebar', b, at.get('Rebar_kg', v * 110), e, 'designed: ' + at['Rebar'] if 'Rebar' in at else '110 kg/m3')
        put('exc', b, 0.5 * 0.5 * L, e, '0.50x0.50 x L'); put('lean', b, 0.3 * 0.05 * L, e, '0.30 x 0.05 x L')
        put('backfill', b, 0.5 * 0.5 * L - 0.2 * 0.4 * L, e, 'trench - beam')
    elif name.endswith('-SUB'):
        conc = [p for p in parts if p[-1] == 'concrete']
        tim = [p for p in parts if p[-1] != 'concrete']
        put('c_pier', b, sum(vol(p) for p in conc) + len(conc) * 0.4 * 0.4 * 0.15, e, 'piers + 0.40x0.40x0.15 pads')
        put('exc', b, len(conc) * 0.4 * 0.4 * 0.6, e, '0.40x0.40x0.60 per pier'); put('fw', b, len(conc) * 4 * 0.2 * 0.5, e, 'pier sides')
        put('rebar', b, sum(vol(p) for p in conc) * 40, e, '40 kg/m3')
        put('tm_beam', b, sum(vol(p) for p in tim) * FT3, e, 'bearers m3 x 35.31')
    elif tag == 'S-Slab':
        area = sum((p[2] - p[1]) * (p[4] - p[3]) if p[0] == 'box' else poly_area3(p[1]) for p in parts)
        perim = sum(2 * ((p[2] - p[1]) + (p[4] - p[3])) for p in parts if p[0] == 'box') * 0.8 or 2 * (mm.DA + mm.DB)
        put('c_sl', b, v, e, 'model volume'); put('sand_fill', b, area * 0.10, e, 'area x 0.10'); put('pe', b, area, e, 'area')
        put('mesh', b, area, e, 'area'); put('fw_sl', b, perim * (0.15 if name != 'D-SLAB' else 0.12), e, 'edge perimeter x depth')
        if name == 'D-SLAB':
            put('fl_trowel', b, area, e, 'area')
    elif ifc == 'IfcColumn' and 'timber_dark' in m:
        put('tm_post', b, v * FT3, e, 'm3 x 35.31'); put('tm_treat', b, v * FT3, e, 'm3 x 35.31')
        put('tm_steel', b, v * 25, e, '25 kg/m3')
    elif tag in ('S-Beam', 'S-RoofTimber') and ifc in ('IfcBeam', 'IfcMember'):
        if name.startswith(('RB-', 'RA-', 'PU-', 'D-MONITOR')):
            put('tm_roof', b, v * FT3, e, 'm3 x 35.31')
        else:
            put('tm_beam', b, v * FT3, e, 'm3 x 35.31')
        put('tm_treat', b, v * FT3, e, 'm3 x 35.31'); put('tm_steel', b, v * 25, e, '25 kg/m3')
    # ---------------- roof
    elif ifc == 'IfcRoof':
        p = parts[0]
        pts = p[1]
        put('rf_sheet', b, at['SlopeArea_m2'], e, 'slope area (laps in waste)')
        put('rf_fascia', b, math.dist(pts[0], pts[1]) + math.dist(pts[1], pts[2]) + math.dist(pts[3], pts[0]), e, 'eave + 2 verges')
        put('drip', b, math.dist(pts[0], pts[1]), e, 'eave length')
    # ---------------- walls / screens
    elif ifc == 'IfcWall' and name == 'D-GABION':
        put('gabion', b, v, e, 'model volume')
    elif ifc == 'IfcWall' and mark == 'GB':
        area = sum(poly_area3(p[1]) if p[0] == 'poly' else (p[2] - p[1]) * (p[6] - p[5]) for p in parts)
        if 'plaster' in m:
            put('wl_aac', b, area, e, 'gable area'); put('plaster', b, area * 2, e, '2 faces'); put('paint_wall', b, area * 2, e, '2 faces')
        else:
            put('cl_vboard', b, area, e, 'gable area'); put('stain_ext', b, area, e, 'exterior face')
    elif ifc == 'IfcWall':
        area = at['NetArea_m2(one face)']
        if 'plaster' in m:
            put('wl_aac', b, area, e, 'net area'); put('plaster', b, area * 2, e, '2 faces'); put('paint_wall', b, area * 2, e, '2 faces')
        else:
            put('wl_frame', b, area, e, 'net area'); put('cl_board', b, area * 2, e, 'net area x 2 faces')
            put('stain_ext', b, area * (1 if at.get('IsExternal') else 0), e, 'exterior face')
    elif tag == 'A-Screen':
        bars = [p for p in parts if p[0] == 'bar']
        if name == 'A-SCREEN-W':
            area = 5.7 * 2.0
        else:
            area = sum(bar_len(p) for p in bars if abs(p[1][2] - p[2][2]) > 0.5) * 0.14
        put('screen', b, area, e, 'slat length x spacing'); put('stain_ext', b, area * 2, e, '2 faces')
    # ---------------- floors, decks, ceilings, site
    elif tag == 'A-Deck' and ifc == 'IfcSlab':
        put('deck', b, at['Area_m2'], e, 'model area'); put('deck_oil', b, at['Area_m2'], e, 'model area')
    elif ifc == 'IfcStair':
        put('steps', b, 1, e, 'count')
    elif ifc == 'IfcCovering':
        it = at.get('Item', 'ceil_timber')
        put(it, b, at.get('Area_m2', 7.0 * 3.1), e, 'model area')
    elif name == 'A-TERRACE':
        put('terrace', b, poly_area3(parts[0][1]), e, 'model area')
    elif name == 'D-DRIVE':
        put('gravel', b, poly_area3(parts[0][1]), e, 'model area')
    elif ifc == 'IfcFurniture' and name in ('C-BENCH-N', 'D-BENCH'):
        put('bench', b, 1, e, 'count')
    # ---------------- doors / windows
    elif ifc in ('IfcDoor', 'IfcWindow'):
        w, h = (float(x) for x in at['Size_WxH_m'].split('x'))
        if mark == 'GD1':
            put('al_bifold', b, w * h, e, 'W x H')
        elif mark == 'SD1':
            put('al_slide', b, w * h, e, 'W x H')
        elif ifc == 'IfcDoor':
            put('dr_reclaim', b, 1, e, 'count (schedule %s)' % mark)
        else:
            put('wn_reclaim_l' if w * h > 1.5 else 'wn_reclaim_s', b, 1, e, 'count (schedule %s)' % mark)
    # ---------------- sanitary / MEP
    elif ifc == 'IfcSanitaryTerminal':
        it = at.get('Item') or {'WC': 'san_wc', 'LB': 'san_lav', 'BT': 'san_tub'}[mark]
        put(it, b, at.get('Count', 1), e, 'count')
    elif 'Item' in at and tag[:2] in ('P-', 'E-'):
        it = at['Item']
        if 'Length_m' in at:              # runs: split by building of each segment
            mult = at.get('Count', 1)
            if name == 'P-HOT':
                for bb in ('A', 'C', 'C'):
                    put(it, bb, 3.0, e, '3 m per heater')
                continue
            for p in parts:
                put(it, locate(centroid(p)), bar_len(p) * mult, e, 'segment length' + (' x %d' % mult if mult > 1 else ''))
        else:
            if ifc == 'IfcLightFixture':
                for p in parts:
                    put('e_lpt', locate(centroid(p)), 1, e, 'one point per fixture')
            if name == 'E-POLE':
                put(it, b, 1, e, 'count')
            else:
                for p in parts:
                    put(it, locate(centroid(p)), 1, e, 'count')
put('wp', 'A', 5.25 * 1.3, basis='bath floor + 0.30 upstand'); put('fl_tile', 'A', 5.25, basis='bath floor')
put('wp', 'C', (6.63 + 5.43) * 1.3, basis='bath floors + upstand'); put('fl_tile', 'C', 6.63 + 5.43, basis='bath floors')
put('wt_tile', 'A', 6.0, basis='shower walls 1.80 h'); put('wt_tile', 'C', 12.0, basis='2 showers x 6 m2')
for bb, slab in (('A', 'A-SLAB'), ('C', 'C-SLAB')):
    e = next(x for x in E if x[1] == slab)
    area = sum((p[2] - p[1]) * (p[4] - p[3]) for p in e[5])
    put('fl_polish', bb, area - (5.25 if bb == 'A' else 12.06), e, 'slab area - bathrooms')
for e in E:                                # ridge caps / flashing from ridge beams
    if e[1].startswith('RB-'):
        L = bar_len(e[5][0])
        put('rf_ridge' if len(e[5]) == 1 else 'rf_flash', elem_bld(e), L * (1 if len(e[5]) == 1 else 2), e, 'ridge length')
put('rf_flash', 'C', 10.4, basis='C-W / C-E roof step (valley flashing)')
tim_ext = sum(q for (k, bb), q in Q.items() if k in ('tm_post', 'tm_beam')) / FT3 * 16
put('stain_ext', 'S', tim_ext, basis='exposed posts/beams ~16 m2 per m3')
if DESIGN.get('timber'):
    for bb, n in (('A', 12), ('C', 16)):
        put('wl_brace', bb, n, basis='racking check: 2 diagonal braces per wall line (design_calc.py)')
put('clean', 'S', 1, basis='lump sum'); put('asbuilt', 'S', 1, basis='lump sum')
put('p_test', 'S', 1, basis='lump sum'); put('e_test', 'S', 1, basis='lump sum')

# ------------------------------------------------------------------ BOQ rows
ROWS = []                                  # dicts in output order
for cat, ctitle in CATS:
    for key, it in ITEMS.items():
        if it[0] != cat:
            continue
        for b in ('A', 'C', 'D', 'S'):
            q = Q.get((key, b), 0)
            if q <= 0:
                continue
            unit = it[3]
            qr = round(q, 0) if unit in ('ชุด', 'จุด', 'ก้อน', 'บ่อ', 'งาน') else (round(q, 0) if unit == 'กก.' else round(q, 2))
            ROWS.append(dict(cat=cat, group=it[1], key=key, desc=it[2], bld=b, qty=qr, unit=unit, waste=it[6] / 100,
                             mat=it[4], lab=it[5], act=act_of(it[7], b)))
for r in ROWS:
    r['mat_amt'] = r['qty'] * (1 + r['waste']) * r['mat']
    r['lab_amt'] = r['qty'] * r['lab']
    r['total'] = r['mat_amt'] + r['lab_amt']
DIRECT = sum(r['total'] for r in ROWS)
ACT_COST = defaultdict(float)
for r in ROWS:
    ACT_COST[r['act']] += r['total']

# ------------------------------------------------------------------ calendar + CPM
def workdays(start, n):
    out, d = [], start
    while len(out) < n:
        if d.weekday() < 6 and d not in HOLIDAYS:
            out.append(d)
        d += dt.timedelta(days=1)
    return out


CAL = workdays(START, 400)
dur = {a[0]: a[2] for a in ACTS}
ES, EF = {}, {}
for code, nm, d, preds in ACTS:
    es = 0
    for p, typ, lag in preds:
        es = max(es, (EF[p] if typ == 'FS' else ES[p]) + lag)
    ES[code], EF[code] = es, es + d
FIN = max(EF.values())
LS, LF = {}, {}
for code, nm, d, preds in reversed(ACTS):
    lf = FIN
    for c2, _, _, p2 in ACTS:
        for p, typ, lag in p2:
            if p == code:
                lf = min(lf, LS[c2] - lag if typ == 'FS' else LS[c2] - lag + d)
    LF[code], LS[code] = lf, lf - d
PLAN = []
for code, nm, d, preds in ACTS:
    PLAN.append(dict(code=code, name=nm, dur=d, es=ES[code], ef=EF[code], start=CAL[ES[code]], finish=CAL[EF[code] - 1],
                     tf=LS[code] - ES[code], crit=LS[code] == ES[code], cost=ACT_COST.get(code, 0.0),
                     preds=', '.join(p + ('' if t == 'FS' and not lag else '(%s%+d)' % (t, lag)) for p, t, lag in preds)))
missing = [c for c in ACT_COST if c not in ACT_IDS]
assert not missing, missing
WEEK0 = START - dt.timedelta(days=START.weekday())
NW = (CAL[FIN - 1] - WEEK0).days // 7 + 1
WEEKS = [WEEK0 + dt.timedelta(days=7 * i) for i in range(NW)]
for a in PLAN:                              # fraction of the activity's workdays falling in each week
    days = CAL[a['es']:a['ef']]
    a['wk'] = [sum(1 for x in days if WEEKS[i] <= x < WEEKS[i] + dt.timedelta(days=7)) / len(days) for i in range(NW)]

MILESTONES = [
    (1, 'ฐานรากทุกอาคาร ท่อใต้พื้น และพื้นคอนกรีต A, C เสร็จ', ['PRE', 'F-A', 'F-C', 'F-D', 'F-S', 'UG-A', 'UG-C', 'SL-A', 'SL-C']),
    (2, 'โครงสร้างไม้ A, C, D และซุ้มเสาเสร็จ', ['TS-A', 'TS-C', 'TS-D', 'PG']),
    (3, 'หลังคาทุกอาคารมุงเสร็จ', ['RF-A', 'RF-C', 'RF-D']),
    (4, 'ผนัง ระแนง กำแพงกาเบียน และงานระบบในผนังเสร็จ', ['WL-A', 'WL-C', 'WL-D', 'MR-A', 'MR-C']),
    (5, 'ชานไม้ และงานระบบภายนอก (ถังบำบัด ปั๊ม สายเมน) เสร็จ', ['DK', 'UT']),
    (6, 'ฝ้า พื้นผิว ประตู-หน้าต่างเสร็จ', ['CL-A', 'CL-C', 'FL-A', 'FL-C', 'DW-A', 'DW-C']),
    (7, 'สี สุขภัณฑ์ ดวงโคม เครื่องใช้ไฟฟ้า ติดตั้งเสร็จ', ['PT-A', 'PT-C', 'FX-A', 'FX-C']),
    (8, 'งานภายนอก ทดสอบระบบ ส่งมอบงาน (งวดสุดท้าย)', ['EX', 'TC', 'HO']),
]
assert sorted(c for _, _, cs in MILESTONES for c in cs) == sorted(ACT_IDS)
MS_OF = {c: n for n, _, cs in MILESTONES for c in cs}

# ------------------------------------------------------------------ Excel
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

FONT = 'Tahoma'
f_n = Font(name=FONT, size=10)
f_b = Font(name=FONT, size=10, bold=True)
f_h = Font(name=FONT, size=13, bold=True)
f_in = Font(name=FONT, size=10, color='0000FF')
f_link = Font(name=FONT, size=10, color='008000')
f_w = Font(name=FONT, size=10, bold=True, color='FFFFFF')
fill_h = PatternFill('solid', fgColor='35524A')
fill_cat = PatternFill('solid', fgColor='DCE6E1')
fill_sub = PatternFill('solid', fgColor='EEF3F0')
fill_key = PatternFill('solid', fgColor='FFFF00')
fill_plan = PatternFill('solid', fgColor='8FB3A6')
fill_crit = PatternFill('solid', fgColor='C0504D')
thin = Side(style='thin', color='B7C4BE')
box_b = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM = '#,##0.00;(#,##0.00);-'
NUM0 = '#,##0;(#,##0);-'
PCT = '0.0%;(0.0%);-'


def hdr(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c)
        x.font, x.fill, x.border = f_w, fill_h, box_b
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[CL(i)].width = w


def style_row(ws, row, ncol, font=f_n, fill=None, fmt=None):
    for c in range(1, ncol + 1):
        x = ws.cell(row=row, column=c)
        x.font, x.border = font, box_b
        if fill:
            x.fill = fill
        if fmt and c in fmt:
            x.number_format = fmt[c]


wb = Workbook()
# ---------------- สรุป (ปร.5)
ws = wb.active
ws.title = 'สรุป'
ws['A1'] = 'แบบสรุปราคาค่าก่อสร้าง (แบบ ปร.5 อย่างย่อ)'; ws['A1'].font = f_h
ws['A2'] = 'โครงการ: ' + PROJECT; ws['A2'].font = f_n
ws['A3'] = 'ปริมาณถอดจากโมเดล BIM (make_model.py) - ราคาต่อหน่วยเป็นค่าประมาณ ปี 2569 - ต้องตรวจสอบก่อนใช้ประกวดราคา'; ws['A3'].font = f_n
hdr(ws, 5, ['ลำดับ', 'รายการ', 'ค่าวัสดุ (บาท)', 'ค่าแรงงาน (บาท)', 'รวมค่างานต้นทุน (บาท)', 'สัดส่วน'], [8, 46, 18, 18, 20, 10])
for i, (cat, ctitle) in enumerate(CATS):
    r = 6 + i
    ws.cell(row=r, column=1, value=i + 1)
    ws.cell(row=r, column=2, value=ctitle)
    ws.cell(row=r, column=3, value="=SUMIFS(BOQ!$I$4:$I$1000,BOQ!$O$4:$O$1000,\"%s\")" % cat)
    ws.cell(row=r, column=4, value="=SUMIFS(BOQ!$K$4:$K$1000,BOQ!$O$4:$O$1000,\"%s\")" % cat)
    ws.cell(row=r, column=5, value='=C%d+D%d' % (r, r))
    ws.cell(row=r, column=6, value='=IF($E$10=0,0,E%d/$E$10)' % r)
    style_row(ws, r, 6, fmt={3: NUM, 4: NUM, 5: NUM, 6: PCT})
    for c in (3, 4):
        ws.cell(row=r, column=c).font = f_link
ws.cell(row=10, column=2, value='รวมค่างานต้นทุน (Direct cost)')
for c, col in ((3, 'C'), (4, 'D'), (5, 'E')):
    ws.cell(row=10, column=c, value='=SUM(%s6:%s9)' % (col, col))
ws.cell(row=10, column=6, value='=SUM(F6:F9)')
style_row(ws, 10, 6, f_b, fill_cat, {3: NUM, 4: NUM, 5: NUM, 6: PCT})
ws.cell(row=11, column=2, value='Factor F (ค่าดำเนินการ กำไร ดอกเบี้ย ภาษีมูลค่าเพิ่ม)')
ws.cell(row=11, column=5, value=FACTOR_F)
ws.cell(row=11, column=5).comment = Comment('ASSUMED: same Factor F as the TYPE03 template (1.2846). Look up the Comptroller-General '
                                            'Factor F table for the real budget size, advance payment, retention, interest and VAT.', 'boq.py')
style_row(ws, 11, 6, fmt={5: '0.0000'})
ws.cell(row=11, column=5).font, ws.cell(row=11, column=5).fill = f_in, fill_key
ws.cell(row=12, column=2, value='ค่าก่อสร้างรวม (Direct cost x Factor F)')
ws.cell(row=12, column=5, value='=E10*E11')
style_row(ws, 12, 6, f_b, fill_cat, {5: NUM})
ws.cell(row=13, column=2, value='ค่าก่อสร้างรวม (ปัดเศษหลักพัน)')
ws.cell(row=13, column=5, value='=ROUNDDOWN(E12,-3)')
style_row(ws, 13, 6, f_b, fmt={5: NUM0})
area_bld = round(foot, 2)
ws.cell(row=15, column=2, value='พื้นที่อาคาร (พื้นคอนกรีต A + C + D) ตร.ม.')
ws.cell(row=15, column=5, value=area_bld); style_row(ws, 15, 6, fmt={5: NUM}); ws.cell(row=15, column=5).font = f_in
ws.cell(row=15, column=5).comment = Comment('From the model: A-SLAB + C-SLAB + D-SLAB plan areas (decks not included).', 'boq.py')
ws.cell(row=16, column=2, value='ราคาต่อ ตร.ม. พื้นที่อาคาร (รวม Factor F)')
ws.cell(row=16, column=5, value='=IF(E15=0,0,E12/E15)'); style_row(ws, 16, 6, fmt={5: NUM})
ws.cell(row=18, column=2, value='ระยะเวลาก่อสร้าง (วันทำงาน)')
ws.cell(row=18, column=5, value="=MAX('แผนงาน'!G4:G100)"); style_row(ws, 18, 6, fmt={5: '0'}); ws.cell(row=18, column=5).font = f_link
ws.cell(row=19, column=2, value='วันเริ่มงาน / วันแล้วเสร็จตามแผน')
ws.cell(row=19, column=3, value=START); ws.cell(row=19, column=5, value="=MAX('แผนงาน'!I4:I100)")
style_row(ws, 19, 6, fmt={3: 'dd/mm/yyyy', 5: 'dd/mm/yyyy'})
ws.cell(row=19, column=3).font, ws.cell(row=19, column=3).fill = f_in, fill_key
ws.cell(row=19, column=3).comment = Comment('ASSUMED start date. Dates in the plan sheet are computed by boq.py (working days Mon-Sat '
                                            'without the listed holidays) - rerun boq.py after changing it.', 'boq.py')
ws.cell(row=21, column=2, value='ตัวอักษรสีน้ำเงิน/พื้นเหลือง = ค่าที่แก้ได้ ; สีเขียว = ดึงจากชีตอื่น ; ราคาต่อหน่วยแก้ในชีต BOQ คอลัมน์ G และ J')
ws.cell(row=21, column=2).font = f_n
ch = BarChart(); ch.type = 'bar'; ch.style = 10; ch.title = 'ค่างานต้นทุนแยกหมวด (บาท)'; ch.legend = None
ch.add_data(Reference(ws, min_col=5, min_row=6, max_row=9)); ch.set_categories(Reference(ws, min_col=2, min_row=6, max_row=9))
ch.height, ch.width = 7, 16
ws.add_chart(ch, 'H5')

# ---------------- BOQ (ปร.4)
wsb = wb.create_sheet('BOQ')
wsb['A1'] = 'บัญชีแสดงปริมาณวัสดุและราคา (แบบ ปร.4) - ' + PROJECT; wsb['A1'].font = f_h
cols = ['ลำดับ', 'รายการ', 'อาคาร', 'ปริมาณ', 'หน่วย', 'เผื่อ', 'ราคาวัสดุ/หน่วย', 'จำนวนเงินค่าวัสดุ', 'ค่าวัสดุรวมเผื่อ', 'ค่าแรง/หน่วย',
        'จำนวนเงินค่าแรง', 'รวมค่าวัสดุและค่าแรง', 'กิจกรรม', 'งวด', 'หมวด', 'ที่มาของปริมาณ']
hdr(wsb, 3, cols, [7, 58, 9, 11, 7, 7, 12, 14, 14, 11, 14, 16, 8, 6, 6, 34])
wsb.freeze_panes = 'C4'
r = 4
ROW_OF = []
cat_rows = []
for cat, ctitle in CATS:
    wsb.cell(row=r, column=1, value=cat); wsb.cell(row=r, column=2, value=ctitle)
    style_row(wsb, r, 16, f_b, fill_cat)
    start = r + 1
    r += 1
    grp = None
    n = 0
    for row in [x for x in ROWS if x['cat'] == cat]:
        if row['group'] != grp:
            grp = row['group']
            wsb.cell(row=r, column=2, value=grp)
            style_row(wsb, r, 16, f_b, fill_sub)
            r += 1
        n += 1
        vals = [n, row['desc'], BLD[row['bld']], row['qty'], row['unit'], row['waste'], row['mat'], '=D{0}*G{0}'.format(r),
                '=D{0}*(1+F{0})*G{0}'.format(r), row['lab'], '=D{0}*J{0}'.format(r), '=I{0}+K{0}'.format(r), row['act'],
                MS_OF[row['act']], cat, 'ถอดจากโมเดล (ดูชีต ถอดปริมาณ: %s)' % row['key']]
        for c, v in enumerate(vals, 1):
            wsb.cell(row=r, column=c, value=v)
        style_row(wsb, r, 16, fmt={4: '#,##0.00', 6: '0%', 7: NUM, 8: NUM, 9: NUM, 10: NUM, 11: NUM, 12: NUM})
        for c in (4, 6, 7, 10):
            wsb.cell(row=r, column=c).font = f_in
        row['xl_row'] = r
        r += 1
    wsb.cell(row=r, column=2, value='รวม ' + ctitle)
    for c, col in ((8, 'H'), (9, 'I'), (11, 'K'), (12, 'L')):
        wsb.cell(row=r, column=c, value='=SUMIFS({0}{1}:{0}{2},$O{1}:$O{2},"{3}")'.format(col, start, r - 1, cat))
    style_row(wsb, r, 16, f_b, fill_cat, {8: NUM, 9: NUM, 11: NUM, 12: NUM})
    cat_rows.append(r)
    r += 2
wsb.cell(row=r, column=2, value='รวมค่างานต้นทุนทั้งหมด')
wsb.cell(row=r, column=12, value='=' + '+'.join('L%d' % x for x in cat_rows))
style_row(wsb, r, 16, f_b, fill_cat, {12: NUM})
wsb.cell(row=r + 2, column=2, value='หมายเหตุ: ปริมาณ (ตัวน้ำเงิน) ถอดจากโมเดล - ราคาวัสดุ/ค่าแรง (ตัวน้ำเงิน) เป็นราคาประมาณ ปี 2569 '
         'ไม่รวม Factor F ; "เผื่อ" คิดเฉพาะค่าวัสดุ ; ไม้เก่าคิดเป็น ลบ.ฟ. (1 ลบ.ม. = 35.31 ลบ.ฟ.)').font = f_n

# ---------------- ถอดปริมาณ
wst = wb.create_sheet('ถอดปริมาณ')
hdr(wst, 1, ['องค์ประกอบ (BIM)', 'IFC', 'Tag', 'อาคาร', 'รหัสรายการ', 'รายการ BOQ', 'ปริมาณ', 'หน่วย', 'วิธีคิด'], [24, 22, 15, 9, 13, 50, 11, 8, 36])
wst.freeze_panes = 'A2'
for i, t in enumerate(TRACE, 2):
    el, ifc, tag, b, item, q, unit, basis = t
    for c, v in enumerate([el, ifc, tag, BLD[b], item, ITEMS[item][2], q, unit, basis], 1):
        wst.cell(row=i, column=c, value=v)
    style_row(wst, i, 9, fmt={7: '#,##0.000'})
wst.auto_filter.ref = 'A1:I%d' % (len(TRACE) + 1)

# ---------------- ประตูหน้าต่าง
wsd = wb.create_sheet('ประตูหน้าต่าง')
hdr(wsd, 1, ['รหัส', 'ประเภท', 'กว้าง (ม.)', 'สูง (ม.)', 'จำนวนที่มี (ชิ้น)', 'ติดตั้งในโมเดล (ชิ้น)', 'สำรอง (ชิ้น)', 'ตำแหน่ง'],
    [8, 10, 10, 10, 13, 15, 11, 60])
for i, m in enumerate(sorted(mm.SCHED), 2):
    h, w, n = mm.SCHED[m]
    used = mm.PLACED.get(m, [])
    for c, v in enumerate([m, 'ประตู' if m[0] == 'D' else 'หน้าต่าง', w, h, n, len(used), '=E%d-F%d' % (i, i), ', '.join(used) or '(สำรองทั้งหมด)'], 1):
        wsd.cell(row=i, column=c, value=v)
    style_row(wsd, i, 8, fmt={3: '0.00', 4: '0.00'})
nr = len(mm.SCHED) + 2
wsd.cell(row=nr, column=2, value='รวม %d รหัส' % len(mm.SCHED))
for c, col in ((5, 'E'), (6, 'F'), (7, 'G')):
    wsd.cell(row=nr, column=c, value='=SUM(%s2:%s%d)' % (col, col, nr - 1))
style_row(wsd, nr, 8, f_b, fill_cat)
wsd.cell(row=nr + 2, column=2, value='นับเป็นชิ้นจริง: W03, W04, W11 มีรหัสละ 3 ชิ้น และ W05, W06 มีรหัสละ 2 ชิ้น ตามตารางในแบบ ("3 PIECES") ส่วนรหัสอื่นมีรหัสละ 1 ชิ้น').font = f_n

# ---------------- แผนงาน (CPM + Gantt by week)
wsp = wb.create_sheet('แผนงาน')
pc = ['รหัส', 'กิจกรรม', 'ระยะเวลา (วันทำงาน)', 'กิจกรรมก่อนหน้า', 'งวด', 'ES', 'EF', 'วันเริ่ม', 'วันเสร็จ', 'Float', 'วิกฤต',
      'มูลค่างานต้นทุน (บาท)', 'น้ำหนักงาน']
hdr(wsp, 3, pc + [w.strftime('%d/%m') for w in WEEKS], [8, 46, 10, 20, 6, 6, 6, 11, 11, 7, 7, 16, 9] + [4.2] * NW)
wsp['A1'] = 'แผนงานก่อสร้าง (CPM) - วันทำงาน จ.-ส. ไม่รวมวันหยุดนักขัตฤกษ์ที่ระบุ - คอลัมน์สัปดาห์ = วันจันทร์ต้นสัปดาห์'; wsp['A1'].font = f_b
wsp.freeze_panes = 'C4'
NP = len(PLAN)
for i, a in enumerate(PLAN):
    r = 4 + i
    vals = [a['code'], a['name'], a['dur'], a['preds'], MS_OF[a['code']], a['es'], a['ef'], a['start'], a['finish'], a['tf'],
            'ใช่' if a['crit'] else '', '=SUMIFS(BOQ!$L$4:$L$1000,BOQ!$M$4:$M$1000,A%d)' % r, '=IF($L$%d=0,0,L%d/$L$%d)' % (4 + NP, r, 4 + NP)]
    for c, v in enumerate(vals, 1):
        wsp.cell(row=r, column=c, value=v)
    style_row(wsp, r, len(pc), fmt={8: 'dd/mm/yy', 9: 'dd/mm/yy', 12: NUM, 13: PCT})
    wsp.cell(row=r, column=12).font = f_link
    for wi, frac in enumerate(a['wk']):
        x = wsp.cell(row=r, column=len(pc) + 1 + wi)
        x.border = box_b
        if frac > 0:
            x.fill = fill_crit if a['crit'] else fill_plan
a['xl_end'] = r
wsp.cell(row=4 + NP, column=2, value='รวม')
wsp.cell(row=4 + NP, column=12, value='=SUM(L4:L%d)' % (3 + NP)); wsp.cell(row=4 + NP, column=13, value='=SUM(M4:M%d)' % (3 + NP))
style_row(wsp, 4 + NP, len(pc), f_b, fill_cat, {12: NUM, 13: PCT})
wsp.cell(row=6 + NP, column=2, value='แถบแดง = กิจกรรมวิกฤต (Float = 0), แถบเขียว = มีเวลาลอยตัว ; SS+n = เริ่มหลังกิจกรรมก่อนหน้าเริ่ม n วัน').font = f_n

# ---------------- S-Curve (weekly planned value from activity cost x fraction of workdays per week)
wsw = wb.create_sheet('การกระจายงาน')
wsw['A1'] = 'สัดส่วนวันทำงานของแต่ละกิจกรรมในแต่ละสัปดาห์ (ใช้คำนวณ S-Curve)'; wsw['A1'].font = f_b
hdr(wsw, 2, ['รหัส', 'มูลค่า'] + ['W%d' % (i + 1) for i in range(NW)], [8, 14] + [6] * NW)
for i, a in enumerate(PLAN):
    r = 3 + i
    wsw.cell(row=r, column=1, value=a['code'])
    wsw.cell(row=r, column=2, value="='แผนงาน'!L%d" % (4 + i)).font = f_link
    for wi, frac in enumerate(a['wk']):
        wsw.cell(row=r, column=3 + wi, value=round(frac, 6)).number_format = '0.00'
wss = wb.create_sheet('S-Curve')
hdr(wss, 1, ['สัปดาห์', 'เริ่มสัปดาห์', 'มูลค่างานตามแผน (บาท)', 'สะสม (บาท)', 'สะสม (%)', 'ผลงานจริงสะสม (%) - กรอก'], [9, 12, 20, 18, 11, 22])
for wi in range(NW):
    r = 2 + wi
    col = CL(3 + wi)
    wss.cell(row=r, column=1, value='W%d' % (wi + 1))
    wss.cell(row=r, column=2, value=WEEKS[wi])
    wss.cell(row=r, column=3, value="=SUMPRODUCT('การกระจายงาน'!$B$3:$B${0},'การกระจายงาน'!{1}$3:{1}${0})".format(2 + NP, col))
    wss.cell(row=r, column=4, value='=SUM($C$2:C%d)' % r)
    wss.cell(row=r, column=5, value='=IF(SUM($C$2:$C$%d)=0,0,D%d/SUM($C$2:$C$%d))' % (1 + NW, r, 1 + NW))
    style_row(wss, r, 6, fmt={2: 'dd/mm/yy', 3: NUM, 4: NUM, 5: PCT, 6: PCT})
    wss.cell(row=r, column=6).font, wss.cell(row=r, column=6).fill = f_in, fill_key
lc = LineChart(); lc.title = 'S-Curve แผนงาน vs ผลงานจริง'; lc.y_axis.title = 'สะสม %'; lc.x_axis.title = 'สัปดาห์'
lc.add_data(Reference(wss, min_col=5, min_row=1, max_row=1 + NW), titles_from_data=True)
lc.add_data(Reference(wss, min_col=6, min_row=1, max_row=1 + NW), titles_from_data=True)
lc.set_categories(Reference(wss, min_col=1, min_row=2, max_row=1 + NW)); lc.y_axis.number_format = '0%'
lc.height, lc.width = 9, 20
wss.add_chart(lc, 'H2')

# ---------------- งวดงาน
wsm = wb.create_sheet('งวดงาน')
hdr(wsm, 1, ['งวด', 'เงื่อนไขการเบิก (งานที่ต้องแล้วเสร็จ)', 'กิจกรรม', 'กำหนดเสร็จตามแผน', 'มูลค่าต้นทุน (บาท)', 'มูลค่ารวม Factor F (บาท)',
             'สัดส่วน', 'สะสม'], [6, 52, 34, 14, 18, 20, 9, 9])
for i, (n, txt, cs) in enumerate(MILESTONES):
    r = 2 + i
    fin = max(a['finish'] for a in PLAN if a['code'] in cs)
    wsm.cell(row=r, column=1, value=n); wsm.cell(row=r, column=2, value=txt); wsm.cell(row=r, column=3, value=', '.join(cs))
    wsm.cell(row=r, column=4, value=fin)
    wsm.cell(row=r, column=5, value='=SUMIFS(BOQ!$L$4:$L$1000,BOQ!$N$4:$N$1000,A%d)' % r)
    wsm.cell(row=r, column=6, value="=E%d*'สรุป'!$E$11" % r)
    wsm.cell(row=r, column=7, value='=IF(SUM($E$2:$E$9)=0,0,E%d/SUM($E$2:$E$9))' % r)
    wsm.cell(row=r, column=8, value='=SUM($G$2:G%d)' % r)
    style_row(wsm, r, 8, fmt={4: 'dd/mm/yy', 5: NUM, 6: NUM, 7: PCT, 8: PCT})
wsm.cell(row=10, column=2, value='รวม')
wsm.cell(row=10, column=5, value='=SUM(E2:E9)'); wsm.cell(row=10, column=6, value='=SUM(F2:F9)'); wsm.cell(row=10, column=7, value='=SUM(G2:G9)')
style_row(wsm, 10, 8, f_b, fill_cat, {5: NUM, 6: NUM, 7: PCT})
wsm.cell(row=12, column=2, value='งวดคิดตามมูลค่าจริงของกิจกรรมในงวด (cost-loaded) - ปรับเป็น % คงที่ตามสัญญาได้ ; หักเงินประกันผลงาน 5% ตามสัญญา').font = f_n

# ---------------- ข้อสมมติ
wsa = wb.create_sheet('ข้อสมมติ')
wsa.column_dimensions['A'].width = 120
notes = [
    'ข้อสมมติและที่มาของตัวเลข',
    '1. ตำแหน่ง/ขนาดในผัง: วัดจากผังบริเวณที่มีสเกล 10 ม. (24.6 px/ม.) ความคลาดเคลื่อน ±0.2 ม. ; ความสูงและความชันหลังคาประมาณจากภาพถ่าย',
    '2. ขนาดประตู-หน้าต่างไม้เก่า: จากตารางในแบบ (D01-D20, W01-W12) ; ตำแหน่งติดตั้งประมาณจากผัง/ภาพ axonometric',
    '3. โครงสร้าง (ฐานราก 0.80x0.80 ตอม่อ 0.20 คานคอดิน 0.20x0.40 เสาไม้ 0.15-0.20) เป็นขนาดสมมติ ยังไม่ได้ออกแบบ - ต้องมีวิศวกรโยธาออกแบบและรับรอง',
    '4. ฐานราก ตอม่อ คานคอดิน: ขนาดและเหล็กเสริมจากรายการคำนวณเบื้องต้น design/design_report.pdf (ACI 318-19, qa = 100 kPa สมมติ) ; ตอม่อชานยังคิด 40 กก./ลบ.ม. ; เหล็กประกับไม้ 25 กก./ลบ.ม.ไม้',
    '4.1 สายป้อน ตู้ไฟ เบรกเกอร์: จากการคำนวณโหลดไฟฟ้าเบื้องต้น (design_calc.py) ; ราคาสาย NYY ประมาณจาก ตร.มม. รวมของทุกแกน',
    '5. งานระบบประปา-สุขาภิบาลและไฟฟ้า: แนวท่อ/สาย ตำแหน่งอุปกรณ์ ขนาดสาย เป็นแบบร่างเพื่อประมาณราคา - ต้องออกแบบโดยวิศวกรเครื่องกล/ไฟฟ้า (มาตรฐาน วสท.)',
    '6. แหล่งไฟฟ้า/น้ำ: สมมติเข้าจากถนนด้านทิศใต้ใกล้ทางรถ (PEA 3 เฟส 15(45)A, ประปาภูมิภาค 1/2") - ค่าธรรมเนียมจริงขึ้นกับการไฟฟ้า/การประปา',
    '7. ราคาต่อหน่วย: ราคาประมาณ ปี 2569 (ฐานเดียวกับแม่แบบ TYPE03 ใน House BIM Studio) ยังไม่รวม Factor F ; ไม้เก่าราคาผันผวนตามแหล่ง',
    '8. Factor F = %.4f (ค่าเดียวกับ TYPE03) - ต้องเปิดตาราง Factor F กรมบัญชีกลางตามวงเงินและเงื่อนไขสัญญาจริง' % FACTOR_F,
    '9. แผนงาน: เริ่ม %s, ทำงาน จ.-ส., หยุดนักขัตฤกษ์ที่ใส่ไว้: %s (ยังไม่รวมวันหยุดตามจันทรคติ)' % (
        START.strftime('%d/%m/%Y'), ', '.join(d.strftime('%d/%m/%y') for d in sorted(HOLIDAYS))),
    '10. ทีมงาน: ทีมฐานราก 1 ทีม, ทีมช่างไม้ 1 ทีม (ทำ A -> C -> D ต่อเนื่อง), ทีมผนัง/ตกแต่งทำ A และ C ซ้อนกัน 7 วัน, ทีมระบบแยก',
    '11. ไม่รวม: เฟอร์นิเจอร์ลอยตัว, ภูมิทัศน์/ต้นไม้ปลูกใหม่, สระน้ำ, รั้วรอบที่ดิน, ระบบ CCTV/อินเทอร์เน็ต, ค่าออกแบบและขออนุญาต',
    '12. ภาพ Gemini_Generated_Image_*.png ในไฟล์ RAR เป็นบ้านคนละหลัง (หลังคาแบน) ไม่ได้นำมาคิด',
]
for i, t in enumerate(notes, 1):
    wsa.cell(row=i, column=1, value=t).font = f_b if i == 1 else f_n
    wsa.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical='top')

for w in wb.worksheets:
    w.sheet_view.zoomScale = 90
XLSX = os.path.join(HERE, 'BaanSaoYongHin_BOQ_Plan.xlsx')
wb.save(XLSX)

# ------------------------------------------------------------------ tracker data
os.makedirs(os.path.join(HERE, 'tracker'), exist_ok=True)
cat_tot = {c: sum(r['total'] for r in ROWS if r['cat'] == c) for c, _ in CATS}
data = {
    'project': PROJECT, 'start': START.isoformat(), 'factorF': FACTOR_F, 'direct': round(DIRECT, 2),
    'holidays': sorted(d.isoformat() for d in HOLIDAYS),
    'cats': [{'code': c, 'title': t, 'total': round(cat_tot[c], 2)} for c, t in CATS],
    'acts': [{'code': a['code'], 'name': a['name'], 'dur': a['dur'], 'start': a['start'].isoformat(), 'finish': a['finish'].isoformat(),
              'es': a['es'], 'ef': a['ef'], 'tf': a['tf'], 'crit': a['crit'], 'cost': round(a['cost'], 2), 'ms': MS_OF[a['code']],
              'preds': a['preds'], 'bld': (a['code'].split('-')[1] if '-' in a['code'] and a['code'].split('-')[1] in 'ACD' else 'S')}
             for a in PLAN],
    'weeks': [w.isoformat() for w in WEEKS],
    'planWeekly': [round(sum(a['cost'] * a['wk'][i] for a in PLAN), 2) for i in range(NW)],
    'milestones': [{'no': n, 'text': t, 'acts': cs, 'value': round(sum(ACT_COST[c] for c in cs), 2),
                    'due': max(a['finish'] for a in PLAN if a['code'] in cs).isoformat()} for n, t, cs in MILESTONES],
    'boq': [{'cat': r['cat'], 'group': r['group'], 'desc': r['desc'], 'bld': r['bld'], 'qty': r['qty'], 'unit': r['unit'],
             'total': round(r['total'], 2), 'act': r['act']} for r in ROWS],
    'elemAct': ELEM_ACT,
}
with open(os.path.join(HERE, 'tracker', 'plan.js'), 'w', encoding='utf-8') as f:
    f.write('window.PLAN = ')
    json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
    f.write(';\n')

print('BOQ rows:', len(ROWS), ' direct cost: {:,.2f}'.format(DIRECT), ' x F: {:,.2f}'.format(DIRECT * FACTOR_F))
for c, t in CATS:
    print('  {:<44} {:>14,.2f}'.format(t, cat_tot[c]))
print('activities:', len(PLAN), ' duration:', FIN, 'workdays ', PLAN[0]['start'], '->', max(a['finish'] for a in PLAN),
      ' critical:', ' '.join(a['code'] for a in PLAN if a['crit']))
print('elements linked to activities (4D):', len(ELEM_ACT), '/', len(E))
